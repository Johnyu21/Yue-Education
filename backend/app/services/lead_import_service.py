"""Bounded Excel preview and atomic employee-confirmed import."""
import hashlib
from pathlib import Path
from uuid import uuid5
from openpyxl import load_workbook
from app.core.errors import InvalidRequestError,ForbiddenError,StateConflictError
from app.schemas.lead import LeadManualIn
from app.services.assessment_service import _require_resume_path
from app.services.crm_service import create_lead
from app.core.idempotency import idempotent
FIELDS={'name','mobile','gender','city','education','date_of_birth','study_purpose','interest_level'}


def preview(db,user,data):
    if not user.is_employee or user.employee_role=='NEWCOMER':raise ForbiddenError()
    file=_require_resume_path(user,data.object_key)
    if file.suffix.lower()!='.xlsx':raise InvalidRequestError('批量导入目前支持xlsx')
    if not {'name','mobile'}<=data.mapping.keys() or not data.mapping.keys()<=FIELDS:raise InvalidRequestError('映射必须包含name/mobile且只能使用允许的客户字段')
    book=load_workbook(file,read_only=True,data_only=True)
    try:
        sheet=book[data.sheet_name] if data.sheet_name else book.active
        headers=[str(v or '').strip() for v in next(sheet.iter_rows(min_row=1,max_row=1,max_col=50,values_only=True))]
        if any(headers.count(v)!=1 for v in data.mapping.values()):raise InvalidRequestError('映射表头不存在或不唯一')
        columns={k:headers.index(v) for k,v in data.mapping.items()};rows=[];errors=[]
        for number,row in enumerate(sheet.iter_rows(min_row=2,max_row=202,max_col=50,values_only=True),start=2):
            if all(v is None for v in row):continue
            if number>201:raise InvalidRequestError('一次最多导入200行')
            values={k:row[i] for k,i in columns.items() if row[i] is not None}
            for k in ('name','mobile','education','city','gender','study_purpose','interest_level'):
                if k in values:values[k]=str(values[k]).strip()
            try:
                parsed=LeadManualIn(**values,request_key=uuid5(data.request_key,str(number)))
                from app.core.mobile import normalize_mobile
                parsed.mobile=normalize_mobile(parsed.mobile)
                rows.append({'row':number,'data':parsed.model_dump(mode='json')})
            except Exception:errors.append({'row':number,'message':'字段格式不正确，请检查姓名、手机号、日期及字段长度'})
        if not rows and not errors:raise InvalidRequestError('工作表没有可导入的数据')
        return {'fingerprint':hashlib.sha256(file.read_bytes()).hexdigest(),'rows':rows,'errors':errors,'canApply':bool(rows) and not errors}
    except KeyError as exc:raise InvalidRequestError('工作表不存在') from exc
    finally:book.close()


@idempotent('lead.import')
def apply(db,user,data):
    if not data.confirmed:raise InvalidRequestError('请核对预览后再确认')
    result=preview(db,user,data)
    if result['fingerprint']!=data.fingerprint:raise StateConflictError('文件已变化，请重新预览')
    if result['errors']:raise InvalidRequestError('请修正所有行错误后重试')
    results=[{'row':row['row'],**create_lead(db,user,LeadManualIn(**row['data']))} for row in result['rows']]
    return {'items':results,'count':len(results)}
