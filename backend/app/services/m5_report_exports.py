from io import BytesIO
from html import escape
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.pagesizes import A4

def export_docx(report):
    doc=Document()
    section=doc.sections[0];section.top_margin=Cm(2.3);section.bottom_margin=Cm(2.3)
    section.left_margin=Cm(2.5);section.right_margin=Cm(2.5)
    for name in ["Normal","Title","Heading 1","Heading 2"]:
        style=doc.styles[name]
        style.font.name="Microsoft YaHei"
        style.font.color.rgb=RGBColor(0,0,0)
        style.element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
    doc.styles["Normal"].font.size=Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after=Pt(7)
    doc.add_paragraph(report.title,style="Title")
    doc.add_paragraph(str(report.start_date)+" 至 "+str(report.end_date)+"  |  "+report.scope)
    doc.add_paragraph("本报告整理所选范围的工作与指标，供员工回顾及主管核对。数据口径与限制请参见数据说明。")
    doc.add_heading("业务指标",level=1)
    for m in report.metrics:
        doc.add_paragraph(str(m["label"])+"："+str(m["value"])+" "+m["unit"])
    doc.add_heading("报告正文",level=1)
    for line in report.narrative.splitlines():
        if line.strip():
            doc.add_paragraph(line.lstrip("# ").strip())
    if report.warnings:
        doc.add_heading("数据说明",level=1)
        for w in report.warnings:doc.add_paragraph(w)
    out=BytesIO();doc.save(out);return out.getvalue()

def export_pdf(report):
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    body=ParagraphStyle("body",fontName="STSong-Light",fontSize=10.5,leading=18,spaceAfter=8,wordWrap="CJK")
    title=ParagraphStyle("title",parent=body,fontSize=21,leading=29,spaceAfter=16)
    heading=ParagraphStyle("heading",parent=body,fontSize=14,leading=22,spaceBefore=12)
    story=[Paragraph(escape(report.title),title),Paragraph(escape(str(report.start_date)+" 至 "+str(report.end_date)+" | "+report.scope),body),
        Paragraph("本报告供工作回顾与管理核对，数据口径与限制请参见数据说明。",body),Paragraph("业务指标",heading)]
    for m in report.metrics:story.append(Paragraph(escape(str(m["label"])+"："+str(m["value"])+" "+m["unit"]),body))
    story.append(Paragraph("报告正文",heading))
    for line in report.narrative.splitlines():
        if line.strip():story.append(Paragraph(escape(line.lstrip("# ").strip()),body))
    if report.warnings:
        story.append(Paragraph("数据说明",heading))
        story.extend(Paragraph(escape(w),body) for w in report.warnings)
    out=BytesIO();SimpleDocTemplate(out,pagesize=A4,rightMargin=58,leftMargin=58,topMargin=52,bottomMargin=52).build(story)
    return out.getvalue()
