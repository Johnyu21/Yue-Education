"""
简历文本抽取（PDF / DOCX）—— M2

⚠️ 两条底线（对齐 M2 模块文档）：
  1. 抽不到可读文本（如扫描件）→ 422 NO_READABLE_TEXT，
     提示"请提供文字版或 OCR 后重试"，**绝不假装识别完成**。
  2. 抽取结果只用于"字段提取草稿"，不直接改写正式客户事实。

依赖：pypdf / python-docx（可选，见 backend/requirements-m2.txt）；
      未安装时按"无非文本"处理，不抛 500。
"""
from __future__ import annotations

import logging
from pathlib import Path

from app.services.m2_errors import NoReadableTextError

logger = logging.getLogger(__name__)

# 送进 Dify 的简历文本上限（防止 payload 过大）
MAX_TEXT_CHARS = 20000

# 支持直接读取的纯文本扩展名
_PLAIN_TEXT_SUFFIXES = (".txt", ".md")


def _pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:                      # 依赖未装 → 视为无可读文本
        logger.warning("未安装 pypdf，无法解析 PDF")
        raise NoReadableTextError("当前环境不支持 PDF 解析，请提供文字版简历")

    try:
        reader = PdfReader(str(path))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:                 # 损坏/加密/非 PDF
        logger.warning("PDF 解析失败：%s", exc.__class__.__name__)
        raise NoReadableTextError("PDF 无法解析，请提供文字版或 OCR 后重试") from exc

    return "\n".join(pages).strip()


def _docx_text(path: Path) -> str:
    try:
        import docx                        # python-docx
    except ImportError:
        logger.warning("未安装 python-docx，无法解析 DOCX")
        raise NoReadableTextError("当前环境不支持 DOCX 解析，请提供 PDF 或文字版简历")

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        logger.warning("DOCX 解析失败：%s", exc.__class__.__name__)
        raise NoReadableTextError("DOCX 无法解析，请提供文字版简历") from exc

    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:            # 简历常用表格排版
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def extract_text(path: Path) -> str:
    """从文件抽取纯文本；无可读文本抛 422 NO_READABLE_TEXT。"""
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        text = _pdf_text(path)
    elif suffix == ".xlsx":
        text = _xlsx_text(path)
    elif suffix == ".docx":
        text = _docx_text(path)
    elif suffix in _PLAIN_TEXT_SUFFIXES:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
    elif suffix == ".doc":
        raise NoReadableTextError("暂不支持 .doc 旧格式，请另存为 PDF 或 DOCX 后重试")
    else:
        raise NoReadableTextError(f"不支持从 {suffix or '该类型'} 文件提取文本")

    if not text:
        raise NoReadableTextError(
            "未从文件中提取到可读文本（可能是扫描件），请提供文字版或 OCR 后重试")

    if len(text) > MAX_TEXT_CHARS:
        logger.info("简历文本超长，截断至 %s 字符", MAX_TEXT_CHARS)
        text = text[:MAX_TEXT_CHARS]

    return text


def _xlsx_text(path):
    try:
        from openpyxl import load_workbook
        with path.open('rb') as source:
            workbook=load_workbook(source,read_only=True,data_only=True)
            parts=[];size=0
            for sheet in workbook:
                for index,row in enumerate(sheet.iter_rows(max_row=1000,max_col=30,values_only=True),start=1):
                    line=sheet.title+'!R'+str(index)+': '+' | '.join(str(v) for v in row if v is not None)
                    parts.append(line);size+=len(line)
                    if size>=MAX_TEXT_CHARS:break
                if size>=MAX_TEXT_CHARS:break
            workbook.close()
            return '\n'.join(parts)
    except Exception as exc:
        raise NoReadableTextError('Excel无法解析，请检查格式或提供文字版') from exc
